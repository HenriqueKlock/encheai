#!/usr/bin/env python3
"""Gera o arquivo público de postos usado pela aba Postos do app.

Fontes (dados abertos da ANP, públicos):
  - cadastro dos revendedores varejistas (todos os postos autorizados, com endereço e bandeira);
  - Levantamento de Preços de Combustíveis (LPC), planilhas semanais por posto;
  - ações de fiscalização do abastecimento (autos de infração, interdições, notificações…).

Os endereços viram coordenadas pelo Nominatim (OpenStreetMap), no máximo 1 consulta por segundo,
com cache: cada posto é procurado uma vez e o resultado fica guardado na branch `dados`.

Uso:
  python pipeline/postos.py --modo probe                 # só mostra cabeçalhos e amostras das fontes
  python pipeline/postos.py --estado estado --saida estado --uf SC
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import json
import math
import os
import re
import sys
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request

UA = 'Mozilla/5.0 (compatible; caderneta-carro/1.0; +https://github.com/HenriqueKlock/caderneta-carro)'
NOMINATIM_UA = 'caderneta-carro/1.0 (+https://github.com/HenriqueKlock/caderneta-carro)'
LPC_PAGE = 'https://www.gov.br/anp/pt-br/assuntos/precos-e-defesa-da-concorrencia/precos/levantamento-de-precos-de-combustiveis-ultimas-semanas-pesquisadas'
CAD_URL = 'https://www.gov.br/anp/pt-br/centrais-de-conteudo/dados-abertos/arquivos/arquivos-dados-cadastrais-dos-revendedores-varejistas-de-combustiveis-automotivos/dados-cadastrais-revendedores-varejistas-combustiveis-automoveis.csv'
FISC_URL = 'https://www.gov.br/anp/pt-br/centrais-de-conteudo/paineis-dinamicos-da-anp/arquivos-dados-brutos-do-painel-dinamico-da-fiscalizacao-do-abastecimento-da-sfi/dados-fisc-a-partir-2019.xlsx'
NOMINATIM = 'https://nominatim.openstreetmap.org/search'
# Postos do OpenStreetMap (ODbL: pode guardar e redistribuir citando "© colaboradores do OpenStreetMap")
OVERPASS = ['https://overpass-api.de/api/interpreter', 'https://overpass.kumi.systems/api/interpreter']

UF_NOME = {
    'AC': 'Acre', 'AL': 'Alagoas', 'AP': 'Amapá', 'AM': 'Amazonas', 'BA': 'Bahia', 'CE': 'Ceará', 'DF': 'Distrito Federal',
    'ES': 'Espírito Santo', 'GO': 'Goiás', 'MA': 'Maranhão', 'MT': 'Mato Grosso', 'MS': 'Mato Grosso do Sul', 'MG': 'Minas Gerais',
    'PA': 'Pará', 'PB': 'Paraíba', 'PR': 'Paraná', 'PE': 'Pernambuco', 'PI': 'Piauí', 'RJ': 'Rio de Janeiro', 'RN': 'Rio Grande do Norte',
    'RS': 'Rio Grande do Sul', 'RO': 'Rondônia', 'RR': 'Roraima', 'SC': 'Santa Catarina', 'SP': 'São Paulo', 'SE': 'Sergipe', 'TO': 'Tocantins',
}
# Geocodifica primeiro onde o app é mais usado; o resto do estado entra nas rodadas seguintes
PRIORIDADE = {'FLORIANOPOLIS', 'SAOJOSE', 'PALHOCA', 'BIGUACU', 'SANTOAMARODAIMPERATRIZ', 'GOVERNADORCELSORAMOS', 'ANTONIOCARLOS'}
FISC_ANOS = 5
GEO_V = 2
LPC_SEMANAS = 6


def log(*a):
    print(*a, flush=True)


def norm(s) -> str:
    """Maiúsculas, sem acento e só letras/dígitos: casa nomes de coluna e municípios."""
    s = unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode()
    return re.sub(r'[^A-Z0-9]', '', s.upper())


def digits(s) -> str:
    return re.sub(r'\D', '', str(s or ''))


def cnpj_of(s) -> str:
    d = digits(s)
    return d.zfill(14) if 11 < len(d) <= 14 else ''


def title(s) -> str:
    s = re.sub(r'\s+', ' ', str(s or '').strip())
    if not s:
        return ''
    small = {'DA', 'DE', 'DO', 'DAS', 'DOS', 'E'}
    out = []
    for i, w in enumerate(s.split(' ')):
        u = w.upper()
        if re.fullmatch(r'[A-Z]{2}-?\d+', u) or u in {'BR', 'SC', 'LTDA', 'ME', 'EPP', 'S/A', 'SA', 'KM'}:
            out.append(u)
        elif i > 0 and u in small:
            out.append(w.lower())
        else:
            out.append(w[:1].upper() + w[1:].lower())
    return ' '.join(out)


def http_get(url: str, tries: int = 4, ua: str = UA, timeout: int = 180) -> bytes:
    last = None
    for i in range(tries):
        try:
            req = urllib.request.Request(url, headers={'User-Agent': ua, 'Accept': '*/*', 'Accept-Language': 'pt-BR,pt;q=0.9'})
            with urllib.request.urlopen(req, timeout=timeout) as r:
                return r.read()
        except (urllib.error.URLError, TimeoutError, ConnectionError) as e:
            last = e
            log(f'  falhou ({e}); tentando de novo em {5 * (i + 1)}s')
            time.sleep(5 * (i + 1))
    raise RuntimeError(f'não consegui baixar {url}: {last}')


def http_head(url: str) -> dict:
    try:
        req = urllib.request.Request(url, method='HEAD', headers={'User-Agent': UA})
        with urllib.request.urlopen(req, timeout=60) as r:
            return {'lm': r.headers.get('Last-Modified'), 'len': r.headers.get('Content-Length'), 'etag': r.headers.get('ETag')}
    except Exception as e:  # noqa: BLE001
        log('  HEAD falhou:', e)
        return {}


def pick(headers: list[str], *cands: str) -> int:
    """Índice da primeira coluna cujo nome normalizado bate (igual, depois começa com) com um candidato."""
    hn = [norm(h) for h in headers]
    for c in cands:
        if c in hn:
            return hn.index(c)
    for c in cands:
        for i, h in enumerate(hn):
            if h.startswith(c):
                return i
    return -1


def to_date(v) -> str | None:
    if v is None or v == '':
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip()
    m = re.match(r'(\d{1,2})/(\d{1,2})/(\d{4})', s)
    if m:
        return f'{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}'
    m = re.match(r'(\d{4})-(\d{2})-(\d{2})', s)
    if m:
        return m.group(0)
    if re.fullmatch(r'\d+(\.0)?', s):  # número de série do Excel
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(float(s)))).isoformat()
    return None


def to_price(v) -> float | None:
    if v is None or v == '':
        return None
    if isinstance(v, (int, float)):
        x = float(v)
    else:
        s = str(v).strip().replace('R$', '').strip()
        s = s.replace('.', '').replace(',', '.') if ',' in s else s
        try:
            x = float(s)
        except ValueError:
            return None
    return round(x, 3) if 0.5 < x < 30 else None


def product_code(s) -> str | None:
    n = norm(s)
    if 'GLP' in n:
        return None
    if 'GASOLINA' in n:
        return 'ga' if 'ADITIV' in n or 'PREMIUM' in n else 'gc'
    if 'ETANOL' in n or 'ALCOOL' in n:
        return 'et'
    if 'GNV' in n:
        return 'gnv'
    if 'DIESEL' in n:
        return 'd10' if 'S10' in n else 'ds'
    return None


# ---------------------------------------------------------------------------
# Cadastro
# ---------------------------------------------------------------------------
def read_csv_bytes(raw: bytes) -> list[list[str]]:
    for enc in ('utf-8-sig', 'cp1252', 'latin-1'):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:20000]
    try:
        delim = csv.Sniffer().sniff(sample, delimiters=';,|\t').delimiter
    except csv.Error:
        delim = ';' if sample.count(';') > sample.count(',') else ','
    return list(csv.reader(io.StringIO(text), delimiter=delim))


def parse_cadastro(rows: list[list[str]], ufs: set[str]) -> dict[str, dict]:
    hi = next(i for i, r in enumerate(rows[:20]) if any(norm(c).startswith('CNPJ') for c in r))
    h = rows[hi]
    c = {
        'cnpj': pick(h, 'CNPJ'), 'razao': pick(h, 'RAZAOSOCIAL', 'RAZAO'), 'fantasia': pick(h, 'NOMEFANTASIA', 'FANTASIA'),
        'bandeira': pick(h, 'BANDEIRA', 'DISTRIBUIDORA'), 'end': pick(h, 'ENDERECO', 'LOGRADOURO'), 'num': pick(h, 'NUMERO'),
        'bairro': pick(h, 'BAIRRO'), 'cep': pick(h, 'CEP'), 'mun': pick(h, 'MUNICIPIO'), 'uf': pick(h, 'UF', 'ESTADO'),
        'comp': pick(h, 'COMPLEMENTO'),
    }
    log('  colunas do cadastro:', {k: (h[v] if v >= 0 else None) for k, v in c.items()})
    g = lambda r, k: (r[c[k]].strip() if 0 <= c[k] < len(r) else '')  # noqa: E731
    out: dict[str, dict] = {}
    for r in rows[hi + 1:]:
        if not r or g(r, 'uf').upper() not in ufs:
            continue
        cnpj = cnpj_of(g(r, 'cnpj'))
        if not cnpj:
            continue
        end, num = g(r, 'end'), g(r, 'num')
        if c['num'] < 0:
            # O cadastro traz o número junto: "RUA WILLY BARTH,  3845" / "RODOVIA BR 101 KM 252,  S/N"
            m = re.match(r'^(.*?),\s*([^,]*)$', end)
            if m:
                end, num = m.group(1), m.group(2)
        out[cnpj] = {
            'cnpj': cnpj, 'razao': g(r, 'razao'), 'fantasia': g(r, 'fantasia'), 'bandeira': g(r, 'bandeira'),
            'end': end.strip(), 'num': num.strip(), 'comp': g(r, 'comp'), 'bairro': g(r, 'bairro'),
            'cep': digits(g(r, 'cep')), 'mun': g(r, 'mun'), 'uf': g(r, 'uf').upper(),
        }
    return out


# ---------------------------------------------------------------------------
# Preços (LPC)
# ---------------------------------------------------------------------------
def lpc_links(html: str) -> list[tuple[str, str, str]]:
    found = {}
    for href in re.findall(r'href="([^"]*revendas_lpc_[^"]*?\.xlsx)"', html, flags=re.I):
        m = re.search(r'revendas_lpc_(\d{4}-\d{2}-\d{2})_(\d{4}-\d{2}-\d{2})', href)
        if not m:
            continue
        url = urllib.parse.urljoin(LPC_PAGE, href)
        found[m.group(1)] = (m.group(1), m.group(2), url)
    return [found[k] for k in sorted(found, reverse=True)]


def xlsx_rows(raw: bytes):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    for row in ws.iter_rows(values_only=True):
        yield list(row)


def parse_lpc(raw: bytes, ufs: set[str], probe: bool = False, stats: dict | None = None):
    """Preços por posto das UFs pedidas. Com `stats`, junta também os preços de todo o país para as médias."""
    it = xlsx_rows(raw)
    h = None
    for i, r in enumerate(it):
        if any(norm(x) == 'CNPJ' for x in r if x is not None):
            h = [str(x or '') for x in r]
            break
        if i > 30:
            raise RuntimeError('cabeçalho do LPC não encontrado')
    c = {
        'cnpj': pick(h, 'CNPJ'), 'razao': pick(h, 'RAZAO', 'RAZAOSOCIAL'), 'fantasia': pick(h, 'FANTASIA', 'NOMEFANTASIA'),
        'end': pick(h, 'ENDERECO'), 'num': pick(h, 'NUMERO'), 'bairro': pick(h, 'BAIRRO'), 'cep': pick(h, 'CEP'),
        'mun': pick(h, 'MUNICIPIO'), 'uf': pick(h, 'ESTADO', 'UF'), 'bandeira': pick(h, 'BANDEIRA'),
        'prod': pick(h, 'PRODUTO'), 'preco': pick(h, 'PRECODEREVENDA', 'PRECOREVENDA', 'PRECODEVENDA', 'PRECO', 'VALORDEVENDA'),
        'data': pick(h, 'DATADACOLETA', 'DATACOLETA', 'DATA'),
    }
    if probe:
        log('  colunas do LPC:', {k: (h[v] if v >= 0 else None) for k, v in c.items()})
    g = lambda r, k: (r[c[k]] if 0 <= c[k] < len(r) else None)  # noqa: E731
    out = []
    for r in it:
        uf = str(g(r, 'uf') or '').strip().upper()
        uf = uf if len(uf) == 2 else next((s for s, n in UF_NOME.items() if norm(n) == norm(uf)), uf)
        if stats is not None and uf not in ufs:
            code, price = product_code(g(r, 'prod')), to_price(g(r, 'preco'))
            if code and price:
                stats.setdefault(('BR', code), []).append(price)
        if uf not in ufs:
            continue
        cnpj = cnpj_of(g(r, 'cnpj'))
        code = product_code(g(r, 'prod'))
        price = to_price(g(r, 'preco'))
        day = to_date(g(r, 'data'))
        if stats is not None and code and price:
            stats.setdefault(('BR', code), []).append(price)
            stats.setdefault(('UF', code), []).append(price)
            stats.setdefault((norm(g(r, 'mun')), code), []).append(price)
        if not (cnpj and code and price and day):
            continue
        out.append({
            'cnpj': cnpj, 'code': code, 'price': price, 'date': day,
            'info': {
                'cnpj': cnpj, 'razao': str(g(r, 'razao') or ''), 'fantasia': str(g(r, 'fantasia') or ''), 'bandeira': str(g(r, 'bandeira') or ''),
                'end': str(g(r, 'end') or ''), 'num': str(g(r, 'num') or ''), 'comp': '', 'bairro': str(g(r, 'bairro') or ''),
                'cep': digits(g(r, 'cep')), 'mun': str(g(r, 'mun') or ''), 'uf': uf,
            },
        })
    return out


# ---------------------------------------------------------------------------
# Fiscalização
# ---------------------------------------------------------------------------
PROC = {
    'AUTODEINFRACAO': 'I', 'AUTODEINTERDICAO': 'D', 'AUTODEAPREENSAO': 'A', 'ATODEINICIODESUSPENSAO': 'S',
    'NOTIFICACAO': 'N', 'MEDIDAREPARADORADECONDUTA': 'M', 'BOLETIMDEFISCALIZACAO': 'B', 'TERMODECOLETADEAMOSTRA': 'C',
}


def parse_fisc(raw: bytes, ufs: set[str], since: str, probe: bool = False) -> dict[str, dict]:
    it = xlsx_rows(raw)
    h = None
    for i, r in enumerate(it):
        if any(norm(x).startswith('CNPJ') for x in r if x is not None):
            h = [str(x or '') for x in r]
            break
        if i > 30:
            raise RuntimeError('cabeçalho da fiscalização não encontrado')
    c = {
        'uf': pick(h, 'UF'), 'cnpj': pick(h, 'CNPJCPF', 'CNPJ'), 'data': pick(h, 'DATADODF', 'DATADF', 'DATA'),
        'doc': pick(h, 'NUMERODODOCUMENTO', 'NUMERODOCUMENTO', 'NUMERO'), 'proc': pick(h, 'PROCEDIMENTODEFISCALIZACAO', 'PROCEDIMENTO', 'COMPONENTE'),
        'res': pick(h, 'RESULTADO'), 'seg': pick(h, 'SEGMENTOFISCALIZADO', 'SEGMENTO'),
    }
    if probe:
        log('  colunas da fiscalização:', {k: (h[v] if v >= 0 else None) for k, v in c.items()})
    g = lambda r, k: (r[c[k]] if 0 <= c[k] < len(r) else None)  # noqa: E731
    out: dict[str, dict] = {}
    n = kept = 0
    segs: dict[str, int] = {}
    for r in it:
        n += 1
        if str(g(r, 'uf') or '').strip().upper() not in ufs:
            continue
        day = to_date(g(r, 'data'))
        if not day or day < since:
            continue
        seg = norm(g(r, 'seg'))
        if c['seg'] >= 0 and seg and not seg.startswith('REVENDADECOMBUSTIVE'):
            continue  # GLP, TRR, distribuidoras etc. não são postos
        cnpj = cnpj_of(g(r, 'cnpj'))
        if not cnpj:
            continue
        kept += 1
        if probe:
            s = str(g(r, 'seg') or '')
            segs[s] = segs.get(s, 0) + 1
        typ = PROC.get(norm(g(r, 'proc')), 'O')
        st = out.setdefault(cnpj, {'visitas': set(), 'ev': {}})
        st['visitas'].add(day)
        if typ in 'IDASNM':
            txt = re.sub(r'\s+', ' ', str(g(r, 'res') or '')).strip()
            key = (day, typ, txt[:80])
            st['ev'][key] = [day, typ, txt[:180]]
    log(f'  fiscalização: {n} linhas lidas, {kept} da(s) UF(s) desde {since}, {len(out)} CNPJs')
    if probe:
        log('  segmentos:', sorted(segs.items(), key=lambda x: -x[1])[:12])
    return {k: {'n': len(v['visitas']), 'ult': max(v['visitas']), 'ev': sorted(v['ev'].values(), reverse=True)[:12]} for k, v in out.items()}


# ---------------------------------------------------------------------------
# Geocodificação (Nominatim, 1 consulta por segundo, com cache)
# ---------------------------------------------------------------------------
ABREV = [
    (r'^(AV|AVEN)\.?\s', 'AVENIDA '), (r'^R\.?\s', 'RUA '), (r'^ROD\.?\s', 'RODOVIA '), (r'^EST\.?\s', 'ESTRADA '),
    (r'^SERV\.?\s', 'SERVIDAO '), (r'^TV\.?\s', 'TRAVESSA '), (r'^AL\.?\s', 'ALAMEDA '), (r'^PCA\.?\s', 'PRACA '),
    (r'^GEN\.?\s', 'GENERAL '), (r'\bDR\.?\s', 'DOUTOR '), (r'\bPROF\.?\s', 'PROFESSOR '), (r'\bSTA\.?\s', 'SANTA '), (r'\bSTO\.?\s', 'SANTO '),
]


def clean_street(s: str) -> str:
    s = re.sub(r'\s+', ' ', str(s or '').upper()).strip(' ,.-')
    s = re.sub(r',?\s*\bN[º°O]?\.?\s*\d+.*$', '', s)
    for pat, rep in ABREV:
        s = re.sub(pat, rep, s)
    s = re.sub(r'\bSC[\s-]?(\d{3})\b', r'SC-\1', s)
    s = re.sub(r'\bBR[\s-]?(\d{3})\b', r'BR-\1', s)
    return s.strip()


SUFIXOS = r'\s+(LTDA\.?|EIRELI|ME|EPP|S\.?/?A\.?|EM RECUPERACAO JUDICIAL)$'
BANDEIRAS = {'BANDEIRABRANCA': 'Bandeira branca', 'RAIZEN': 'Shell', 'VIBRA': 'Petrobras', 'PETROBRASDISTRIBUIDORASA': 'Petrobras', 'BRANCA': 'Bandeira branca', 'IPIRANGA': 'Ipiranga', 'ALESAT': 'Ale'}


def display_name(p: dict) -> str:
    n = p.get('fantasia') or p.get('razao') or ''
    n = re.sub(r'\s+', ' ', unicodedata.normalize('NFC', n.upper())).strip()
    for _ in range(3):
        n = re.sub(SUFIXOS, '', n).strip(' .-')
    return title(n)[:60]


def display_flag(s: str) -> str:
    k = norm(s)
    for key, val in BANDEIRAS.items():
        if k.startswith(key):
            return val
    return title(s)[:30]


BAIRRO_ABREV = [(r'\bRIB\.?\b', 'RIBEIRAO'), (r'\bSTO\.?\b', 'SANTO'), (r'\bSTA\.?\b', 'SANTA'), (r'\bJD\.?\b', 'JARDIM'),
                (r'\bVL\.?\b', 'VILA'), (r'\bPQ\.?\b', 'PARQUE'), (r'\bCONJ\.?\b', 'CONJUNTO'), (r'\bRES\.?\b', 'RESIDENCIAL'),
                (r'\bN\.? ?S(RA)?\.?\b', 'NOSSA SENHORA'), (r'\bS\.? ?JOAO\b', 'SAO JOAO'), (r'\bLOT\.?\b', 'LOTEAMENTO')]
BAIRRO_TIPOS = {'suburb', 'neighbourhood', 'quarter', 'city_district', 'hamlet', 'village', 'residential', 'locality', 'borough', 'isolated_dwelling', 'town'}


def clean_bairro(s: str) -> str:
    s = re.sub(r'\s+', ' ', str(s or '').upper()).strip(' ,.-')
    if norm(s) in {'', 'CENTRO', 'ZONARURAL', 'INTERIOR', 'RURAL', 'SN', 'NAOINFORMADO'}:
        return '' if norm(s) != 'CENTRO' else 'CENTRO'
    for pat, rep in BAIRRO_ABREV:
        s = re.sub(pat, rep, s)
    return s


def same_name(a: str, b: str) -> bool:
    x, y = norm(a), norm(b)
    if not x or not y:
        return False
    if x in y or y in x:
        return True
    wa = [w for w in re.findall(r'[A-Z]{4,}', unicodedata.normalize('NFKD', a).encode('ascii', 'ignore').decode().upper()) if w not in {'BAIRRO', 'JARDIM', 'VILA', 'PARQUE'}]
    return bool(wa) and wa[0] in y


def addr_key(p: dict) -> str:
    return norm(' '.join([p.get('end', ''), p.get('num', ''), p.get('bairro', ''), p.get('mun', ''), p.get('uf', '')]))


def bbox_km(bb) -> float:
    try:
        s, n, w, e = (float(x) for x in bb)
    except (TypeError, ValueError):
        return 99
    dy = (n - s) * 111.0
    dx = (e - w) * 111.0 * math.cos(math.radians((n + s) / 2))
    return math.hypot(dx, dy)


class SemCota(Exception):
    """Acabou o limite de consultas desta rodada."""


class Geo:
    def __init__(self, cache: dict, budget: int):
        self.cache = cache
        self.budget = budget
        self.calls = 0
        self.last = 0.0

    def _q(self, params: dict) -> list:
        if self.calls >= self.budget:
            raise SemCota
        wait = 1.1 - (time.time() - self.last)
        if wait > 0:
            time.sleep(wait)
        self.last = time.time()
        self.calls += 1
        url = NOMINATIM + '?' + urllib.parse.urlencode({**params, 'format': 'jsonv2', 'addressdetails': 1, 'limit': 3, 'countrycodes': 'br', 'accept-language': 'pt-BR'})
        try:
            return json.loads(http_get(url, tries=3, ua=NOMINATIM_UA, timeout=40))
        except Exception as e:  # noqa: BLE001
            log('   nominatim falhou:', e)
            return []

    @staticmethod
    def _same_city(res: dict, mun: str, uf: str) -> bool:
        a = res.get('address') or {}
        if norm(a.get('state')) != norm(UF_NOME.get(uf, uf)):
            return False
        names = [a.get(k) for k in ('city', 'town', 'village', 'municipality', 'city_district', 'county')]
        return any(norm(x) == norm(mun) for x in names if x)

    def locate(self, p: dict):
        """Devolve (lat, lon, precisão) com precisão 'n' (número), 'r' (rua) ou 'b' (bairro); None se não achou."""
        uf, mun = p['uf'], p['mun']
        street = clean_street(p['end'])
        num = digits(p['num'])[:6]
        base = {'city': title(mun), 'state': UF_NOME.get(uf, uf), 'country': 'Brasil'}
        tries = []
        if street and num and not re.search(r'\bKM\b', street):
            tries.append(('n', {**base, 'street': f'{num} {title(street)}'}))
        if street:
            tries.append(('r', {**base, 'street': title(street)}))
        bairro = clean_bairro(p.get('bairro', ''))
        if bairro:
            tries.append(('b', {'q': f"{title(bairro)}, {title(mun)}, {UF_NOME.get(uf, uf)}, Brasil"}))
        for prec, params in tries:
            for res in self._q(params):
                if not self._same_city(res, mun, uf):
                    continue
                size = bbox_km(res.get('boundingbox'))
                lat, lon = round(float(res['lat']), 5), round(float(res['lon']), 5)
                # Rua comprida (rodovia, avenida que cruza a cidade): o centro dela não serve como endereço
                if prec in ('n', 'r') and res.get('category', res.get('class')) == 'highway' and size > 1.5:
                    break
                # Bairro: só vale se o resultado for mesmo o bairro (não a cidade inteira nem outro lugar)
                if prec == 'b' and (size > 6 or res.get('addresstype') not in BAIRRO_TIPOS or not same_name(res.get('name', ''), bairro)):
                    continue
                a = res.get('address') or {}
                city = next((a[k] for k in ('city', 'town', 'village', 'municipality') if a.get(k)), None)
                return lat, lon, prec, city
        return None


# ---------------------------------------------------------------------------
# OpenStreetMap: nome na fachada e ponto exato dos postos (dados abertos, ODbL)
BIG = {'SHELL': 'Shell', 'RAIZEN': 'Shell', 'IPIRANGA': 'Ipiranga', 'PETROBRAS': 'Petrobras', 'BR': 'Petrobras', 'VIBRA': 'Petrobras', 'ALE': 'Ale', 'ALESAT': 'Ale'}
NOME_COMUM = {'POSTO', 'POSTOS', 'AUTO', 'AUTOPOSTO', 'COMERCIO', 'COMBUSTIVEIS', 'COMBUSTIVEL', 'DERIVADOS', 'PETROLEO', 'LTDA', 'SERVICOS',
              'SERVICO', 'DE', 'DA', 'DO', 'DOS', 'DAS', 'REDE', 'CONVENIENCIA', 'ABASTECIMENTO', 'EIRELI', 'CENTRO', 'GAS', 'STATION'}
TIPO_VIA = {'RUA', 'RODOVIA', 'AVENIDA', 'ALAMEDA', 'TRAVESSA', 'ESTRADA', 'PRACA', 'SERVIDAO', 'ROD', 'AV', 'R', 'VIA', 'ACESSO', 'MARGINAL'}


def words(s) -> list[str]:
    return re.findall(r'[A-Z0-9]+', unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode().upper())


def name_tokens(s) -> set[str]:
    return {w for w in words(s) if len(w) >= 3 and w not in NOME_COMUM and w not in BIG and not w.isdigit()}


def brand_of(s) -> str | None:
    for w in words(s):
        if w in BIG:
            return BIG[w]
    return None


def street_tokens(s) -> set[str]:
    s = clean_street(s)
    hw = {a + b for a, b in re.findall(r'\b(SC|BR)-?(\d{3})\b', s)}
    return hw | {w for w in words(s) if len(w) >= 3 and w not in TIPO_VIA and not w.isdigit() and not re.fullmatch(r'(SC|BR)\d{3}', w)}


def dist_km(a_la, a_lo, b_la, b_lo) -> float:
    r = math.radians
    h = math.sin(r(b_la - a_la) / 2) ** 2 + math.cos(r(a_la)) * math.cos(r(b_la)) * math.sin(r(b_lo - a_lo) / 2) ** 2
    return 2 * 6371 * math.asin(min(1, math.sqrt(h)))


def fetch_osm(uf: str) -> list[dict] | None:
    """Todos os postos (amenity=fuel) do estado no OpenStreetMap, pelo Overpass. None se não respondeu."""
    q = f'[out:json][timeout:240];area["ISO3166-2"="BR-{uf}"]["admin_level"="4"]->.a;nwr["amenity"="fuel"](area.a);out center tags;'
    for url in OVERPASS:
        try:
            data = json.loads(http_get(url + '?' + urllib.parse.urlencode({'data': q}), tries=2, ua=NOMINATIM_UA, timeout=300))
        except Exception as e:  # noqa: BLE001
            log('   overpass falhou:', url, e)
            continue
        out = []
        for e in data.get('elements', []):
            c = e.get('center') or {}
            la, lo = e.get('lat', c.get('lat')), e.get('lon', c.get('lon'))
            if la is None or lo is None:
                continue
            tg = e.get('tags') or {}
            out.append({'id': e.get('type', 'x')[0] + str(e.get('id')), 'la': round(la, 6), 'lo': round(lo, 6),
                        'name': (tg.get('name') or '').strip()[:60], 'brand': (tg.get('brand') or tg.get('operator') or '').strip()[:40],
                        'street': tg.get('addr:street') or '', 'num': tg.get('addr:housenumber') or '', 'city': tg.get('addr:city') or ''})
        return out
    return None


def osm_detail(p: dict, la, lo, prec, o: dict) -> dict:
    """Evidências de que um posto do OpenStreetMap é este posto da ANP (endereço, nome, bandeira, distância)."""
    num = digits(p.get('num'))
    street_ok = len(street_tokens(p.get('end')) & street_tokens(o['street'])) >= 1
    addr = bool(num) and digits(o['num']) == num and street_ok
    d = dist_km(la, lo, o['la'], o['lo']) if la is not None else None
    ov = len(name_tokens((p.get('fantasia') or '') + ' ' + (p.get('razao') or '')) & name_tokens(o['name']))
    return {'d': d, 'addr': addr, 'street': street_ok, 'ov': ov, 'prec': prec}


def osm_score(p: dict, la, lo, prec, o: dict) -> float:
    """O quanto um posto do OpenStreetMap parece ser este posto da ANP (endereço, nome, bandeira, distância)."""
    det = osm_detail(p, la, lo, prec, o)
    addr, street_ok, ov, d = det['addr'], det['street'], det['ov'], det['d']
    if la is None:
        # posto sem coordenada: só pelo endereço completo, na mesma cidade
        return 4.0 if addr and o['city'] and norm(o['city']) == norm(p.get('mun')) else 0
    # sem o endereço completo, ponto aproximado não pode ir longe: redes (Galo, Angeloni…) têm vários postos perto
    if d > (5 if addr else {'n': 0.35, 'r': 0.8, 'b': 1.2}.get(prec, 1.0)):
        return 0
    ob = brand_of(o['brand'] + ' ' + o['name'])
    pb = display_flag(p.get('bandeira', ''))
    pb = pb if pb in ('Shell', 'Ipiranga', 'Petrobras', 'Ale') else None
    s = (3 if addr else 1 if street_ok else 0) + min(ov, 2) * 1.5 + (1 if ob and ob == pb else 0)
    s += 2 if d < 0.05 else 1.2 if d < 0.15 else 0.6 if d < 0.3 else 0
    if ob and pb and ob != pb:
        s -= 2
    evidence = addr or ov > 0 or d < 0.08 or (ob is not None and ob == pb and d < 0.2)
    return s if evidence and s >= 2.5 else 0


def match_osm(cad: dict, geo_cache: dict, osm: list[dict]) -> dict[str, tuple[dict, float]]:
    """Liga cada posto do OpenStreetMap a no máximo um posto da ANP (os pares mais parecidos primeiro)."""
    grid: dict[tuple[int, int], list[dict]] = {}
    for o in osm:
        grid.setdefault((int(o['la'] // 0.05), int(o['lo'] // 0.05)), []).append(o)
    by_addr: dict[tuple[str, str], list[dict]] = {}
    for o in osm:
        if o['num'] and o['city']:
            by_addr.setdefault((norm(o['city']), digits(o['num'])), []).append(o)
    pairs = []
    for cnpj, p in cad.items():
        g = geo_cache.get(cnpj) or {}
        if g.get('p') is not None and g.get('k') == addr_key(p):
            la, lo, prec = g['la'], g['lo'], g['p']
            gy, gx = int(la // 0.05), int(lo // 0.05)
            cands = [o for dy in (-1, 0, 1) for dx in (-1, 0, 1) for o in grid.get((gy + dy, gx + dx), [])]
        else:
            la = lo = prec = None
            cands = by_addr.get((norm(p.get('mun')), digits(p.get('num'))), [])
        for o in cands:
            s = osm_score(p, la, lo, prec, o)
            if s:
                pairs.append((s, cnpj, o))
    pairs.sort(key=lambda x: -x[0])
    used, out = set(), {}
    for s, cnpj, o in pairs:
        if cnpj in out or o['id'] in used:
            continue
        used.add(o['id'])
        out[cnpj] = (o, s)
    return out


# ---------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--modo', default='completo', choices=['completo', 'probe'])
    ap.add_argument('--uf', default='SC')
    ap.add_argument('--estado', default='estado', help='pasta com o cache (branch dados)')
    ap.add_argument('--saida', default=None, help='pasta do JSON público (padrão: a mesma do estado)')
    ap.add_argument('--max-geo', type=int, default=2500, help='consultas ao Nominatim nesta rodada')
    ap.add_argument('--cad', help='arquivo local do cadastro (testes)')
    ap.add_argument('--lpc', nargs='*', help='planilhas LPC locais (testes)')
    ap.add_argument('--fisc', help='planilha local da fiscalização (testes)')
    ap.add_argument('--sem-geo', action='store_true', help='não consulta o Nominatim (testes)')
    ap.add_argument('--osm', default='usar', choices=['usar', 'relatorio', 'nao'],
                    help='OpenStreetMap: usar (nome e ponto no arquivo), relatorio (só um relatório para conferir) ou nao')
    ap.add_argument('--osm-arquivo', help='postos do OpenStreetMap em arquivo local (testes)')
    ap.add_argument('--hoje', default=None)
    a = ap.parse_args()
    ufs = {u.strip().upper() for u in a.uf.split(',') if u.strip()}
    probe = a.modo == 'probe'
    today = dt.date.fromisoformat(a.hoje) if a.hoje else dt.date.today()
    since = (today - dt.timedelta(days=365 * FISC_ANOS)).isoformat()
    est = a.estado
    out_dir = a.saida or est
    os.makedirs(est, exist_ok=True)
    os.makedirs(out_dir, exist_ok=True)

    def load(name, default):
        try:
            with open(os.path.join(est, name), encoding='utf-8') as f:
                return json.load(f)
        except (FileNotFoundError, json.JSONDecodeError):
            return default

    def save(name, obj, folder=None):
        path = os.path.join(folder or est, name)
        with open(path, 'w', encoding='utf-8') as f:
            json.dump(obj, f, ensure_ascii=False, separators=(',', ':'), sort_keys=True)
        log(f'  gravado {path} ({os.path.getsize(path) // 1024} KB)')

    # 1) Cadastro
    log('Cadastro de revendedores…')
    raw = open(a.cad, 'rb').read() if a.cad else http_get(CAD_URL)
    rows = read_csv_bytes(raw)
    if probe:
        log('  linhas:', len(rows), '| cabeçalho:', rows[0])
        log('  amostra:', [r for r in rows[1:4000] if any(x.strip().upper() in ufs for x in r)][:2])
    cad = parse_cadastro(rows, ufs)
    log(f'  {len(cad)} postos em {",".join(sorted(ufs))}')

    # 2) Preços (últimas semanas)
    log('Levantamento de preços…')
    prices: dict[str, dict[str, list]] = {}
    weeks = []
    week_stats: dict = {}
    if a.lpc is not None:
        files = [(f'local{i}', f'local{i}', p) for i, p in enumerate(a.lpc)]
    else:
        files = lpc_links(http_get(LPC_PAGE).decode('utf-8', 'ignore'))[:LPC_SEMANAS]
    log('  semanas:', [f[0] + '→' + f[1] for f in files])
    for ini, fim, url in files:
        try:
            data = open(url, 'rb').read() if a.lpc is not None else http_get(url)
            # médias de preço (Brasil, estado e cada cidade) só da semana mais recente
            st = {} if not weeks else None
            recs = parse_lpc(data, ufs, probe=probe and not weeks, stats=st)
            if st is not None:
                week_stats = st
        except Exception as e:  # noqa: BLE001
            log(f'  semana {ini} ignorada: {e}')
            continue
        weeks.append([ini, fim])
        log(f'  {ini}: {len(recs)} preços em {len({r["cnpj"] for r in recs})} postos')
        for r in recs:
            cur = prices.setdefault(r['cnpj'], {}).get(r['code'])
            if not cur or r['date'] > cur[1]:
                prices[r['cnpj']][r['code']] = [r['price'], r['date']]
            if r['cnpj'] not in cad:
                cad[r['cnpj']] = r['info']
            elif not cad[r['cnpj']].get('fantasia') and r['info']['fantasia']:
                cad[r['cnpj']]['fantasia'] = r['info']['fantasia']
        if probe:
            break

    # 3) Fiscalização (mensal; só baixa de novo se o arquivo mudou)
    log('Fiscalização…')
    fisc_cache = load(f'fisc-{"-".join(sorted(ufs))}.json', {})
    if a.fisc:
        fisc = parse_fisc(open(a.fisc, 'rb').read(), ufs, since, probe)
        fisc_meta = {'local': a.fisc}
    else:
        head = http_head(FISC_URL)
        if not probe and fisc_cache.get('meta') and head and fisc_cache['meta'].get('head') == head and fisc_cache.get('since', '') <= since:
            log('  arquivo não mudou; usando o resumo guardado')
            fisc, fisc_meta = fisc_cache['postos'], fisc_cache['meta']
        else:
            t0 = time.time()
            data = http_get(FISC_URL, timeout=600)
            log(f'  baixado {len(data) // 1024 // 1024} MB em {time.time() - t0:.0f}s')
            fisc = parse_fisc(data, ufs, since, probe)
            fisc_meta = {'head': head, 'baixado': today.isoformat()}
    if probe:
        some = list(fisc.items())[:3]
        log('  amostra fiscalização:', json.dumps(some, ensure_ascii=False)[:1500])
        log('Probe concluído.')
        return
    save(f'fisc-{"-".join(sorted(ufs))}.json', {'meta': fisc_meta, 'since': since, 'postos': fisc})
    fisc_last = max((v['ult'] for v in fisc.values()), default=None)

    # 4) Coordenadas
    log('Coordenadas…')
    geo_cache = load('geocache.json', {})
    geo = Geo(geo_cache, 0 if a.sem_geo else a.max_geo)
    retry_before = (today - dt.timedelta(days=90)).isoformat()

    def need(p):
        g = geo_cache.get(p['cnpj'])
        if not g or g.get('k') != addr_key(p):
            return True
        # v2: bairro só aceito se o resultado for o próprio bairro; rodovia comprida não vale mais como ponto
        if g.get('v') != GEO_V and g.get('p') in ('b', None):
            return True
        return g.get('p') is None and g.get('em', '') < retry_before

    todo = [p for p in cad.values() if need(p)]
    todo.sort(key=lambda p: (p['cnpj'] not in prices, norm(p['mun']) not in PRIORIDADE, p['mun'], p['cnpj']))
    log(f'  {len(todo)} postos sem coordenada (limite desta rodada: {geo.budget} consultas)')
    done = 0
    try:
        for p in todo:
            r = geo.locate(p)
            geo_cache[p['cnpj']] = {'k': addr_key(p), 'em': today.isoformat(), 'v': GEO_V, **({'la': r[0], 'lo': r[1], 'p': r[2], 'mn': r[3]} if r else {'p': None})}
            done += 1
            if done % 100 == 0:
                log(f'  {done} processados, {geo.calls} consultas')
                save('geocache.json', geo_cache)
    except SemCota:
        log('  limite de consultas desta rodada atingido; o resto fica para a próxima')
    save('geocache.json', geo_cache)

    # 5) OpenStreetMap: nome na fachada e ponto exato (dados abertos, guardados na branch dados)
    osm_match: dict[str, tuple[dict, float]] = {}
    osm_quando = None
    if a.osm != 'nao':
        log('OpenStreetMap…')
        osm_cache = load(f'osm-{"-".join(sorted(ufs))}.json', {})
        osm = None
        if a.osm_arquivo:
            osm = json.load(open(a.osm_arquivo, encoding='utf-8'))
        elif not a.sem_geo:
            osm = fetch_osm(sorted(ufs)[0])
        if osm is not None:
            osm_cache = {'em': today.isoformat(), 'postos': osm}
            save(f'osm-{"-".join(sorted(ufs))}.json', osm_cache)
        osm = osm_cache.get('postos') or []
        osm_quando = osm_cache.get('em')
        osm_match = match_osm(cad, geo_cache, osm) if osm else {}
        log(f'  {len(osm)} postos no OpenStreetMap; {len(osm_match)} ligados a postos da ANP')

    def osm_use(cnpj, g):
        """
        (nome, ponto) do OpenStreetMap para este posto.
        Nome: só com evidência além da vizinhança (endereço, nome parecido, mesma rua bem perto) e só se diz algo
        além da bandeira ("Shell", "Posto Ipiranga" não ajudam). Ponto: só quando o nosso é aproximado ou a
        ligação é forte, e sem saltos grandes (a não ser com endereço completo e nome parecido).
        """
        m = osm_match.get(cnpj)
        if not m:
            return None, None
        o, s = m
        p = cad[cnpj]
        has_pt = g.get('p') is not None
        det = osm_detail(p, g.get('la') if has_pt else None, g.get('lo') if has_pt else None, g.get('p'), o)
        d = det['d']
        name_ok = bool(name_tokens(o['name'])) and (det['addr'] or det['ov'] > 0 or (det['street'] and d is not None and d < 0.1) or (d is not None and d < 0.05))
        if not has_pt:
            point = (o['la'], o['lo'])
        elif d <= {'n': 0.3, 'r': 0.8, 'b': 1.2}.get(g['p'], 0.3) or (det['addr'] and det['ov'] > 0 and d <= 3):
            point = (o['la'], o['lo']) if (g['p'] in ('r', 'b') or s >= 3.5) else None
        else:
            point = None
        return (o['name'] if name_ok else None), point

    if a.osm == 'relatorio' and osm_match:
        rel = []
        for cnpj, (o, s) in sorted(osm_match.items()):
            p = cad[cnpj]
            g = geo_cache.get(cnpj) or {}
            nome, ponto = osm_use(cnpj, g)
            rel.append({'c': cnpj, 'mun': title(p['mun']), 'anp': display_name(p), 'osm': o['name'], 'nome': nome, 'nota': round(s, 2), 'antes': g.get('p'),
                        'move_m': round(dist_km(g['la'], g['lo'], ponto[0], ponto[1]) * 1000) if ponto and g.get('la') is not None else None,
                        'ganha_ponto': g.get('p') is None and ponto is not None})
        save(f'osm-relatorio-{"-".join(sorted(ufs))}.json', {'em': today.isoformat(), 'ligados': len(rel), 'postos': rel})

    # 6) Arquivo público
    postos = []
    for cnpj, p in cad.items():
        g = geo_cache.get(cnpj) or {}
        if g.get('k') != addr_key(p):
            g = {}
        nome_osm, ponto_osm = osm_use(cnpj, g) if a.osm == 'usar' else (None, None)
        if g.get('p') is None and not ponto_osm:
            continue
        num = digits(p['num'])
        endereco = title(clean_street(p['end'])) + (', ' + num if num else '')
        item = {
            'c': cnpj, 'n': display_name(p), 'b': display_flag(p['bandeira']), 'e': endereco[:90],
            'br': title(p['bairro'])[:40], 'm': g.get('mn') if g.get('mn') and norm(g.get('mn')) == norm(p['mun']) else title(p['mun']),
            'la': g.get('la'), 'lo': g.get('lo'), 'pg': g.get('p'),
        }
        if ponto_osm:
            # ponto do OpenStreetMap: é o do próprio posto (melhor que o do endereço)
            item['la'], item['lo'], item['pg'] = ponto_osm[0], ponto_osm[1], 'n'
        if nome_osm:
            item['o'] = nome_osm
        if cnpj in prices:
            item['pr'] = prices[cnpj]
        if cnpj in fisc:
            item['f'] = fisc[cnpj]
        postos.append(item)
    postos.sort(key=lambda x: x['c'])
    sem = len(cad) - len(postos)
    # Preço médio da semana (pesquisa da ANP): referência para todo mundo, no país, no estado e em cada cidade
    def media(k):
        v = week_stats.get(k) or []
        return [round(sum(v) / len(v), 2), len(v)] if len(v) >= 3 else None
    codes = ('gc', 'et', 'ga', 'd10', 'gnv')
    medias = None
    if weeks and week_stats:
        medias = {'semana': weeks[0], 'br': {c: m for c in codes if (m := media(('BR', c)))}, 'uf': {c: m for c in codes if (m := media(('UF', c)))},
                  'mun': {}}
        for (k, c) in week_stats:
            if k not in ('BR', 'UF') and (m := media((k, c))):
                medias['mun'].setdefault(k, {})[c] = m
    doc = {
        'v': 1, 'uf': sorted(ufs), 'gerado': dt.datetime.now(dt.timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ'),
        'fontes': {'lpc_semanas': weeks, 'fisc_desde': since, 'fisc_ate': fisc_last, 'cadastro': today.isoformat(),
                   **({'osm': osm_quando} if a.osm == 'usar' and osm_quando else {})},
        'contagem': {'cadastro': len(cad), 'com_local': len(postos), 'sem_local': sem, 'com_preco': sum(1 for x in postos if 'pr' in x)},
        'postos': postos,
        **({'medias': medias} if medias else {}),
    }
    save(f'postos-{"-".join(sorted(ufs))}.json', doc, out_dir)
    log('Resumo:', json.dumps(doc['contagem']))


if __name__ == '__main__':
    try:
        main()
    except KeyboardInterrupt:
        sys.exit(130)
